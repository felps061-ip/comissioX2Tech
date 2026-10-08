from pathlib import Path

import openpyxl

from converter import (
    _clean_pan_product_name,
    _create_margin_variant,
    _clean_product_base_name,
    _contract_type,
    _detect_agreement,
    _detect_bank,
    _finanto_contract_type,
    _parse_term_range,
    _parse_ticket_range,
    _parse_pan_percentage,
    _parse_pan_rates,
    _pan_contract_type,
    _product_name,
    _read_neo_dataframe,
    _should_create_margin_variant,
    convert_c6_refin,
)


def test_convert_sample_c6(tmp_path: Path) -> None:
    source = Path(__file__).resolve().parents[2] / "outputs" / "c6_2tech_work" / "20261006135632.xls"
    output = tmp_path / "c6.xlsx"

    result = convert_c6_refin(source, output)

    assert result.row_count == 56
    workbook = openpyxl.load_workbook(output, data_only=False)
    sheet = workbook.active
    assert sheet.max_row == 57
    assert sheet.max_column == 41
    assert sheet["C2"].value == "C6BANK"
    assert sheet["D2"].value == "INSS"
    assert sheet["E2"].value == "220077 C6 BANK DIGITAL INSS REFIN 1 (+90D CARENCIA) WEB PLUS 1,85% À 1,85%"
    assert sheet["M2"].value == "FÍSICO"
    assert sheet["M3"].value == "DIGITAL"
    assert sheet["V2"].value == "1,85"
    assert sheet["W2"].value == "1,85"
    assert sheet["X2"].value == "10,17"
    assert sheet["AA2"].value == "4,5765"
    assert sheet["E30"].value == "220077 C6 BANK DIGITAL INSS REFIN + MARGEM 1 (+90D CARENCIA) WEB PLUS 1,85% À 1,85%"
    assert sheet["L30"].value == "REFIN + MARGEM"
    assert not any(
        isinstance(cell.value, str) and cell.value.startswith("=")
        for row in sheet.iter_rows()
        for cell in row
    )


def test_parse_term_range() -> None:
    assert _parse_term_range(108) == (108, 108)
    assert _parse_term_range("108") == (108, 108)
    assert _parse_term_range("1 a 108") == (1, 108)
    assert _parse_term_range("1 até 96") == (1, 96)
    assert _parse_term_range("Prazo 13 - 84") == (13, 84)


def test_parse_ticket_range() -> None:
    assert _parse_ticket_range("C6 BANK NOVO INSS ML TKT MAIOR 20K WEB 108X") == (20000, 999999)
    assert _parse_ticket_range("C6 BANK NOVO INSS ML MAIOR 20K WEB 108X") == (20000, 999999)
    assert _parse_ticket_range("C6 BANK NOVO INSS ML TKT 6K A 20K WEB 108X") == (6000, 20000)
    assert _parse_ticket_range("C6 BANK NOVO INSS ML 6K A 20K WEB 108X") == (6000, 20000)
    assert _parse_ticket_range("C6 BANK NOVO INSS ML TKT ATE 6K WEB 108X") == (0, 6000)
    assert _parse_ticket_range("C6 BANK DIGITAL INSS REFIN WEB 108X") == (0, 999999)
    assert _parse_ticket_range("DIGIO INSS REFIN LIQ TKT 1500 1,80 TAB 8814 DIG 108X") == (1500, 999999)
    assert _parse_ticket_range("DAYCOVAL INSS NOVO TAB TKT 15K DIG 108X") == (15000, 999999)


def test_clean_product_base_name_removes_tx_range() -> None:
    assert (
        _clean_product_base_name("C6 BANK INSS PORT TX 1,850~1,35 WEB 108X")
        == "C6 BANK INSS PORT WEB"
    )


def test_clean_product_base_name_removes_single_tx_rate() -> None:
    assert _clean_product_base_name("C6 BANK SIAPE REFIN PORT TX 1,80 NORMAL WEB 96X") == "C6 BANK SIAPE REFIN PORT NORMAL WEB"


def test_clean_product_base_name_removes_term_range() -> None:
    assert _clean_product_base_name("C6 BANK SIAPE PORT TX 1,80~1,60 NORMAL WEB 1 A 96X") == "C6 BANK SIAPE PORT NORMAL WEB"


def test_detect_agreement() -> None:
    assert _detect_agreement("C6 BANK SIAPE DIGITAL NOVO WEB 96X") == "SIAPE"
    assert _detect_agreement("C6 BANK INSS REFIN WEB 108X") == "INSS"
    assert _detect_agreement("DAYCOVAL FEDERAL REFIN LIQ RECUPERACAO DIG 120X") == "SIAPE"


def test_digio_rules() -> None:
    name = "DIGIO INSS REFIN+MARGEM LIQ TKT 1500 1,80 TAB 8814 DIG 96X"
    assert _detect_bank(name) == "DIGIO"
    assert _contract_type(name, "REFIN") == "REFIN+MARGEM"
    assert _product_name("8814", name, 1.8, 1.8, "DIGIO") == (
        "DIGIO INSS REFIN+MARGEM LIQ TKT 1500 1,80 TAB 8814 DIG"
    )


def test_daycoval_rules() -> None:
    name = "DAYCOVAL INSS REFIN LIQ TAB 2 DIG 108X"
    assert _detect_bank(name) == "DAYCOVAL"
    assert _product_name("803499", name, 1.8, 1.8, "DAYCOVAL") == (
        "803499 DAYCOVAL INSS REFIN LIQ TAB 2 DIG 1,80% À 1,80%"
    )


def test_safra_rules() -> None:
    name = "SAFRA INSS REFIN 1,85 LIQ DIGITAL"
    assert _detect_bank(name) == "SAFRA"
    assert _product_name("321412", name, 1.85, 1.85, "SAFRA") == (
        "321412 SAFRA INSS REFIN 1,85 LIQ DIGITAL"
    )


def test_banrisul_rules() -> None:
    name = "BANRISUL INSS REFIN CART LIQ 347N NORMAL FIGITAL 48X"
    assert _detect_bank(name) == "BANRISUL"
    assert _product_name("347N", name, 1.85, 1.85, "BANRISUL") == (
        "347N BANRISUL INSS REFIN CART LIQ 347N NORMAL FIGITAL 1,85% À 1,85%"
    )


def test_banco_do_brasil_rules() -> None:
    name = "BB MAIS CREDITO NAO CONSIGNADO REFIN LIQ TAB 13 A 96X"
    assert _detect_bank(name) == "BANCO DO BRASIL"
    assert _detect_agreement(name) == "CRÉDITO PESSOAL"
    assert _product_name("", name, 4.75, 5.94, "BANCO DO BRASIL") == (
        "BB MAIS CREDITO NAO CONSIGNADO REFIN LIQ TAB 4,75% À 5,94%"
    )
    assert _detect_agreement("BB MAIS CONSIGNADO DO TRABALHADOR NOVO TAB 13 A 96X") == "PRIVADO/CLT"


def test_finanto_rules() -> None:
    assert _detect_bank("1036 FINANTO INSS NOVO 108X - TAXA 1,85%") == "FINANTO"
    assert _finanto_contract_type("FINANTO INSS NOVO 108X") == "NOVO"
    assert _finanto_contract_type("FINANTO INSS REFIN PRIME 108X") == "REFIN"
    assert _finanto_contract_type("FINANTO INSS REFIN DA PORT PRIME 108X") == "REFIN DA PORT"
    assert _finanto_contract_type("FINANTO INSS PORTABILIDADE 96 A 108X") == "PORTABILIDADE"


def test_create_margin_variant() -> None:
    source = [None] * 41
    source[2] = "SAFRA"
    source[4] = "321412 SAFRA INSS REFIN LIQ DIGITAL"
    source[11] = "REFIN"
    variant = _create_margin_variant(source)
    assert variant[4] == "321412 SAFRA INSS REFIN + MARGEM LIQ DIGITAL"
    assert variant[11] == "REFIN + MARGEM"
    assert _should_create_margin_variant(source)

    source[2] = "DIGIO"
    assert not _should_create_margin_variant(source)

    source[2] = "BANCO DO BRASIL"
    assert not _should_create_margin_variant(source)

    source[2] = "NEO CREDITO"
    assert not _should_create_margin_variant(source)

    source[2] = "FINANTO"
    assert not _should_create_margin_variant(source)


def test_pan_rules() -> None:
    name = "SIAPE_PENS_RFN_NORMAL - Taxa 1,76% a 1,80% - Cód. 501530 - VENDA DIGITAL - VD (Venda Digital)"
    assert _clean_pan_product_name(name) == "SIAPE_PENS_RFN_NORMAL - Taxa 1,76% a 1,80% - Cód. 501530"
    assert _parse_pan_rates(name) == (1.76, 1.8)
    assert _pan_contract_type(name) == "REFIN"
    assert _pan_contract_type("SIAPE_SERV_PRFN_FLEX1 - Taxa 1,71% a 1,75%") == "REFIN DA PORT"
    assert _pan_contract_type("INSS_NOV_NORMAL_A_ - Taxa 1,85%") == "NOVO"
    assert _pan_contract_type("INSS_PAN13_NORMAL - Taxa 1,81% a 1,85%") == "REFIN"
    assert _parse_pan_percentage(0.036) == 3.6
    assert _parse_pan_percentage("3,60%") == 3.6
    assert _parse_pan_percentage("-") is None


def test_pan_margin_variant_names() -> None:
    source = [None] * 41
    source[2] = "PAN"
    source[11] = "REFIN"

    source[4] = "SIAPE_PENS_RFN_NORMAL - Taxa 1,76%"
    assert _create_margin_variant(source)[4] == "SIAPE_PENS_REFIN + MARGEM_NORMAL - Taxa 1,76%"

    source[4] = "SIAPE_SERV_PRFN_FLEX1 - Taxa 1,71%"
    assert _create_margin_variant(source)[4] == "SIAPE_SERV_REFIN + MARGEM DA PORT_FLEX1 - Taxa 1,71%"

    source[4] = "INSS_PAN13_NORMAL - Taxa 1,81%"
    assert _create_margin_variant(source)[4] == "INSS_PAN13_REFIN + MARGEM_NORMAL - Taxa 1,81%"


def test_neo_layout_rules() -> None:
    raw = __import__("pandas").DataFrame(
        [
            [None, "#VALUE!", None, "TABELA COMISSÃO", None, None, None],
            [None, "CONVÊNIO", "TIPO", "CÓDIGO TABELA", "PRAZO", "COEF", "CMS"],
            [None, "SIAPE", "COMPRA DÍVIDA", "NCDT_096-419_428248", 96, 0.0428, 0.27],
        ]
    )
    neo = _read_neo_dataframe(raw)
    row = neo.iloc[0]
    assert row["NEO_CONVENIO"] == "SIAPE"
    assert row["NEO_TIPO"] == "COMPRA DÍVIDA"
    assert row["NEO_CODIGO"] == "NCDT_096-419_428248"
    assert row["NEO_PRAZO"] == 96
    assert row["NEO_CMS"] == 0.27
