from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


OUTPUT_HEADERS = [
    "Id do Produto na Origem",
    "Id do Produto no Seu Sistema",
    "Banco",
    "Convênio",
    "Tabela/Nome do Produto",
    "Código no Banco",
    "Id de Vigência",
    "Início",
    "Fim",
    "Prazo Inicial",
    "Prazo Final",
    "Tipo de Contrato",
    "Tipo de Formalização",
    "Fator",
    "Taxa a.m",
    "Id de Prazo/Faixa",
    "Idade Mínima",
    "Idade Máxima",
    "Valor Contrato Inicial",
    "Valor Contrato Final",
    "Valor Contrato Referência",
    "Taxa Inicial",
    "Taxa Final",
    "À Vista (Empresa)",
    "Bônus (Empresa)",
    "Diferido (Empresa)",
    "À Vista (Repasse 1)",
    "Bônus (Repasse 1)",
    "Diferido (Repasse 1)",
    "À Vista (Repasse 2)",
    "Bônus (Repasse 2)",
    "Diferido (Repasse 2)",
    "À Vista (Repasse 3)",
    "Bônus (Repasse 3)",
    "Diferido (Repasse 3)",
    "À Vista (Repasse 4)",
    "Bônus (Repasse 4)",
    "Diferido (Repasse 4)",
    "À Vista (Repasse 5)",
    "Bônus (Repasse 5)",
    "Diferido (Repasse 5)",
]


@dataclass(frozen=True)
class ConversionResult:
    output_path: Path
    row_count: int


def convert_c6_refin(input_path: str | Path, output_path: str | Path, start_date: str | None = None) -> ConversionResult:
    input_path = Path(input_path)
    output_path = Path(output_path)

    if not input_path.exists():
        raise FileNotFoundError(f"Arquivo de entrada não encontrado: {input_path}")

    source_rows = _read_bevi_html_table(input_path)
    converted_rows = [_convert_row(row, start_date) for row in source_rows]
    if not converted_rows:
        raise ValueError("Nenhuma linha válida de comissão C6 foi encontrada.")

    margin_rows = [_create_margin_variant(row) for row in converted_rows if _should_create_margin_variant(row)]
    converted_rows.extend(margin_rows)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    _write_workbook(converted_rows, output_path)
    return ConversionResult(output_path=output_path, row_count=len(converted_rows))


def _read_bevi_html_table(input_path: Path) -> Iterable[pd.Series]:
    if input_path.suffix.lower() in {".xlsx", ".xlsm"}:
        raw_df = pd.read_excel(input_path, header=None, keep_default_na=False)
        if _is_neo_dataframe(raw_df):
            return _read_neo_dataframe(raw_df).iterrows()
        else:
            df = pd.read_excel(input_path, keep_default_na=False)
    else:
        tables = pd.read_html(input_path, keep_default_na=False)
        if not tables:
            raise ValueError("Não encontrei nenhuma tabela no arquivo informado.")
        df = tables[0]

    if _is_finanto_dataframe(df):
        return df[df["Banco"].astype(str).str.contains("FINANTO", case=False, na=False)].iterrows()
    if _is_pan_dataframe(df):
        # The PAN files interleave commission records with descriptive rows.
        df = df[df["Banco"].astype(str).str.contains("PAN", case=False, na=False)]
        return df.iterrows()

    if df.shape[1] < 12:
        raise ValueError("A tabela BEVI não tem o layout esperado para C6/refino.")

    return df.iterrows()


def _convert_row(index_and_row: tuple[int, pd.Series], start_date: str | None) -> list[object | None]:
    _, row = index_and_row
    if _is_neo_row(row):
        return _convert_neo_row(row, start_date)
    if _is_finanto_row(row):
        return _convert_finanto_row(row, start_date)
    if _is_pan_row(row):
        return _convert_pan_row(row, start_date)

    source_date = _parse_date(start_date or row.iloc[2])
    source_name = _clean_text(row.iloc[3])
    contract_type = _clean_text(row.iloc[4])
    rate_initial = _as_float(row.iloc[5])
    rate_final = _as_float(row.iloc[6])
    term_initial, term_final = _parse_term_range(row.iloc[7])
    value_initial, value_final = _parse_ticket_range(source_name)
    company_upfront = _as_float(row.iloc[8])
    bank = _detect_bank(source_name)
    product_name = _product_name(_clean_text(row.iloc[1]), source_name, rate_initial, rate_final, bank)
    repasse_one = None if bank == "BANCO DO BRASIL" else _format_percentage(company_upfront * 0.45, places=4)

    return [
        None,
        None,
        bank,
        _detect_agreement(source_name),
        product_name,
        None,
        None,
        source_date,
        None,
        term_initial,
        term_final,
        _contract_type(source_name, contract_type),
        _formalization_type(source_name),
        None,
        None,
        None,
        18,
        999,
        value_initial,
        value_final,
        "LÍQUIDO",
        _format_rate(rate_initial),
        _format_rate(rate_final),
        _format_rate(company_upfront),
        None,
        None,
        repasse_one,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
    ]


def _is_pan_dataframe(df: pd.DataFrame) -> bool:
    return {"Banco", "Tabela", "Comissionamento", "Prazo", "Flat Repassada"}.issubset(df.columns)


def _is_pan_row(row: pd.Series) -> bool:
    return {"Banco", "Tabela", "Comissionamento", "Prazo", "Flat Repassada"}.issubset(row.index)


def _is_finanto_dataframe(df: pd.DataFrame) -> bool:
    return _is_pan_dataframe(df) and df["Banco"].astype(str).str.contains("FINANTO", case=False, na=False).any()


def _is_finanto_row(row: pd.Series) -> bool:
    return _is_pan_row(row) and "FINANTO" in _normalize_text(row["Banco"])


def _is_neo_dataframe(df: pd.DataFrame) -> bool:
    return any(
        {"CONVENIO", "TIPO", "CODIGO TABELA", "PRAZO", "CMS"}.issubset(
            {_normalize_text(value) for value in row.tolist()}
        )
        for _, row in df.iterrows()
    )


def _read_neo_dataframe(raw_df: pd.DataFrame) -> pd.DataFrame:
    header_index = next(
        index
        for index, row in raw_df.iterrows()
        if {"CONVENIO", "TIPO", "CODIGO TABELA", "PRAZO", "CMS"}.issubset(
            {_normalize_text(value) for value in row.tolist()}
        )
    )
    neo_df = raw_df.iloc[header_index + 1 :, :7].copy()
    neo_df.columns = [
        "NEO_IGNORE",
        "NEO_CONVENIO",
        "NEO_TIPO",
        "NEO_CODIGO",
        "NEO_PRAZO",
        "NEO_COEF",
        "NEO_CMS",
    ]
    return neo_df[neo_df["NEO_CONVENIO"].astype(str).str.strip().ne("")]


def _is_neo_row(row: pd.Series) -> bool:
    return {"NEO_CONVENIO", "NEO_TIPO", "NEO_CODIGO", "NEO_PRAZO", "NEO_CMS"}.issubset(row.index)


def _convert_neo_row(row: pd.Series, start_date: str | None) -> list[object | None]:
    if not start_date:
        raise ValueError("A tabela do NEO CRÉDITO não possui data de vigência. Informe o início da vigência.")

    agreement = _clean_text(row["NEO_CONVENIO"])
    contract_type = _clean_text(row["NEO_TIPO"])
    table_code = _clean_text(row["NEO_CODIGO"])
    company_upfront = _parse_pan_percentage(row["NEO_CMS"])
    term_initial, term_final = _parse_term_range(row["NEO_PRAZO"])

    return [
        None,
        None,
        "NEO CREDITO",
        agreement,
        f"{agreement} - {contract_type} {table_code}",
        None,
        None,
        _parse_date(start_date),
        None,
        term_initial,
        term_final,
        contract_type,
        "DIGITAL",
        None,
        None,
        None,
        18,
        999,
        0,
        999999,
        "LÍQUIDO",
        None,
        None,
        _format_rate(company_upfront) if company_upfront is not None else None,
        None,
        None,
        _format_percentage(company_upfront * 0.45, places=4) if company_upfront is not None else None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
    ]


def _convert_pan_row(row: pd.Series, start_date: str | None) -> list[object | None]:
    if not start_date:
        raise ValueError("A tabela do Banco PAN não possui data de vigência. Informe o início da vigência.")

    source_name = _clean_text(row["Tabela"])
    rate_initial, rate_final = _parse_pan_rates(source_name)
    company_upfront = _parse_pan_percentage(row["Flat Repassada"])
    contract_type = _pan_contract_type(source_name)

    return [
        None,
        None,
        "PAN",
        _detect_agreement(source_name),
        _clean_pan_product_name(source_name),
        None,
        None,
        _parse_date(start_date),
        None,
        *_parse_term_range(row["Prazo"]),
        contract_type,
        "DIGITAL",
        None,
        None,
        None,
        18,
        999,
        0,
        999999,
        "BRUTO" if "BRUTO" in _normalize_text(row["Comissionamento"]) else "LÍQUIDO",
        _format_rate(rate_initial),
        _format_rate(rate_final),
        _format_rate(company_upfront) if company_upfront is not None else None,
        None,
        None,
        _format_percentage(company_upfront * 0.45, places=4) if company_upfront is not None else None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
    ]


def _convert_finanto_row(row: pd.Series, start_date: str | None) -> list[object | None]:
    if not start_date:
        raise ValueError("A tabela da FINANTO não possui data de vigência. Informe o início da vigência.")

    source_name = _clean_text(row["Tabela"])
    rate_initial, rate_final = _parse_pan_rates(source_name)
    company_upfront = _parse_pan_percentage(row["Flat Repassada"])

    return [
        None,
        None,
        "FINANTO",
        _detect_agreement(source_name),
        _clean_pan_product_name(source_name),
        None,
        None,
        _parse_date(start_date),
        None,
        *_parse_term_range(row["Prazo"]),
        _finanto_contract_type(source_name),
        "DIGITAL",
        None,
        None,
        None,
        18,
        999,
        0,
        999999,
        "BRUTO" if "BRUTO" in _normalize_text(row["Comissionamento"]) else "LÍQUIDO",
        _format_rate(rate_initial),
        _format_rate(rate_final),
        _format_rate(company_upfront) if company_upfront is not None else None,
        None,
        None,
        _format_percentage(company_upfront * 0.45, places=4) if company_upfront is not None else None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
    ]


def _product_name(
    bank_description: str,
    source_name: str,
    rate_initial: float,
    rate_final: float,
    bank: str,
) -> str:
    name_without_term = _clean_product_base_name(source_name)
    if bank == "DIGIO":
        return name_without_term
    if bank == "SAFRA":
        return f"{bank_description} {name_without_term}"
    if bank == "BANCO DO BRASIL":
        return f"{name_without_term} {_format_rate(rate_initial)}% À {_format_rate(rate_final)}%"
    return (
        f"{bank_description} {name_without_term} "
        f"{_format_rate(rate_initial)}% À {_format_rate(rate_final)}%"
    )


def _clean_product_base_name(source_name: str) -> str:
    cleaned = re.sub(
        r"\s+\d+\s*(?:A|ATE|ATÉ|-)\s*\d+\s*X\s*$",
        "",
        source_name,
        flags=re.IGNORECASE,
    ).strip()
    cleaned = re.sub(r"\s+\d+\s*X\s*$", "", cleaned, flags=re.IGNORECASE).strip()
    cleaned = re.sub(
        r"\bTX\s*\d+(?:[,.]\d+)?(?:\s*(?:~|A|À|-)\s*\d+(?:[,.]\d+)?%?)?",
        "",
        cleaned,
        flags=re.IGNORECASE,
    )
    return _clean_text(cleaned)


def _clean_pan_product_name(source_name: str) -> str:
    cleaned = re.sub(
        r"\s*-\s*VENDA\s+DIGITAL\s*-\s*VD\s*\(VENDA\s+DIGITAL\)\s*$",
        "",
        source_name,
        flags=re.IGNORECASE,
    )
    return _clean_text(cleaned)


def _parse_pan_rates(source_name: str) -> tuple[float, float]:
    match = re.search(
        r"\bTAXA\s+(\d+(?:[,.]\d+)?)\s*%?(?:\s*(?:A|À|ATE|ATÉ|-)\s*(\d+(?:[,.]\d+)?)\s*%?)?",
        source_name,
        flags=re.IGNORECASE,
    )
    if not match:
        raise ValueError(f"Taxa não identificada na tabela PAN: {source_name}")
    initial = _as_float(match.group(1))
    final = _as_float(match.group(2) or match.group(1))
    return initial, final


def _parse_pan_percentage(value: object) -> float | None:
    text = _clean_text(value)
    if text in {"", "-"}:
        return None

    numeric_text = text.replace("%", "").strip()
    number = _as_float(numeric_text)
    return number if "%" in text else round(number * 100, 8)


def _pan_contract_type(source_name: str) -> str:
    normalized = _normalize_text(source_name)
    if "PRFN" in normalized:
        return "REFIN DA PORT"
    if "RFN" in normalized or "PAN13" in normalized:
        return "REFIN"
    if "NOV" in normalized:
        return "NOVO"
    if "PORTABILIDADE" in normalized:
        return "PORTABILIDADE"
    raise ValueError(f"Operação PAN não identificada no nome da tabela: {source_name}")


def _finanto_contract_type(source_name: str) -> str:
    normalized = _normalize_text(source_name)
    if "REFIN DA PORT" in normalized:
        return "REFIN DA PORT"
    if "REFIN" in normalized:
        return "REFIN"
    if "NOVO" in normalized:
        return "NOVO"
    if "PORTABILIDADE" in normalized:
        return "PORTABILIDADE"
    raise ValueError(f"Operação FINANTO não identificada no nome da tabela: {source_name}")


def _detect_agreement(source_name: str) -> str:
    normalized = _normalize_text(source_name)
    if "CREDITO NAO CONSIGNADO" in normalized:
        return "CRÉDITO PESSOAL"
    if "CONSIGNADO DO TRABALHADOR" in normalized:
        return "PRIVADO/CLT"
    for agreement in ("INSS", "SIAPE"):
        if agreement in normalized:
            return agreement
    if "DAYCOVAL FEDERAL" in normalized:
        return "SIAPE"
    raise ValueError(f"Convênio C6 não identificado no nome da tabela: {source_name}")


def _detect_bank(source_name: str) -> str:
    normalized = _normalize_text(source_name)
    if "BANCO DO BRASIL" in normalized or re.search(r"\bBB\b", normalized):
        return "BANCO DO BRASIL"
    if "DIGIO" in normalized:
        return "DIGIO"
    if "SAFRA" in normalized:
        return "SAFRA"
    if "DAYCOVAL" in normalized:
        return "DAYCOVAL"
    if "BANRISUL" in normalized:
        return "BANRISUL"
    if "PAN" in normalized:
        return "PAN"
    if "FINANTO" in normalized:
        return "FINANTO"
    if "C6" in normalized:
        return "C6BANK"
    raise ValueError(f"Banco não identificado no nome da tabela: {source_name}")


def _contract_type(source_name: str, source_contract_type: str) -> str:
    if "REFIN+MARGEM" in _normalize_text(source_name):
        return "REFIN+MARGEM"
    return source_contract_type


def _formalization_type(source_name: str) -> str:
    normalized = _normalize_text(source_name)
    if "WEB PLUS" in normalized:
        return "FÍSICO"
    return "DIGITAL"


def _should_create_margin_variant(row: list[object | None]) -> bool:
    bank = _normalize_text(row[2])
    contract_type = _normalize_text(row[11])
    return (
        bank not in {"DIGIO", "BANCO DO BRASIL", "BB", "NEO CREDITO", "FINANTO"}
        and "REFIN" in contract_type
        and "MARGEM" not in contract_type
    )


def _create_margin_variant(row: list[object | None]) -> list[object | None]:
    margin_row = row.copy()
    product_name = str(margin_row[4])
    if re.search(r"\bREFIN\b", product_name, flags=re.IGNORECASE):
        margin_row[4] = re.sub(r"\bREFIN\b", "REFIN + MARGEM", product_name, count=1, flags=re.IGNORECASE)
    elif re.search(r"PRFN", product_name, flags=re.IGNORECASE):
        margin_row[4] = re.sub(r"PRFN", "REFIN + MARGEM DA PORT", product_name, count=1, flags=re.IGNORECASE)
    elif re.search(r"RFN", product_name, flags=re.IGNORECASE):
        margin_row[4] = re.sub(r"RFN", "REFIN + MARGEM", product_name, count=1, flags=re.IGNORECASE)
    elif re.search(r"PAN13", product_name, flags=re.IGNORECASE):
        margin_row[4] = re.sub(r"PAN13", "PAN13_REFIN + MARGEM", product_name, count=1, flags=re.IGNORECASE)
    margin_row[11] = "REFIN + MARGEM"
    return margin_row


def _clean_text(value: object) -> str:
    return re.sub(r"\s+", " ", str(value).strip())


def _as_int(value: object) -> int:
    return int(float(str(value).replace(",", ".")))


def _parse_term_range(value: object) -> tuple[int, int]:
    text = _clean_text(value).lower()
    numbers = re.findall(r"\d+", text)
    if len(numbers) >= 2:
        return int(numbers[0]), int(numbers[1])
    if len(numbers) == 1:
        term = int(numbers[0])
        return term, term
    raise ValueError(f"Prazo inválido: {value}")


def _parse_ticket_range(source_name: str) -> tuple[int, int]:
    normalized = _normalize_text(source_name)
    default = (0, 999999)

    range_match = re.search(
        r"(?:\bTKT\s*)?(\d+(?:[,.]\d+)?)\s*K\s*(?:A|ATE|ATÉ|-)\s*(\d+(?:[,.]\d+)?)\s*K",
        normalized,
    )
    if range_match:
        return _money_token_to_number(range_match.group(1), has_k=True), _money_token_to_number(
            range_match.group(2), has_k=True
        )

    max_match = re.search(r"(?:\bTKT\s*)?(?:ATE|ATÉ)\s*(\d+(?:[,.]\d+)?)\s*K", normalized)
    if max_match:
        return 0, _money_token_to_number(max_match.group(1), has_k=True)

    min_match = re.search(r"(?:\bTKT\s*)?(?:MAIOR|ACIMA)\s*(?:DE\s*)?(\d+(?:[,.]\d+)?)\s*K", normalized)
    if min_match:
        return _money_token_to_number(min_match.group(1), has_k=True), default[1]

    digio_ticket_match = re.search(r"\bTKT\s+(\d+(?:[,.]\d+)?)(?!\s*K)\b", normalized)
    if digio_ticket_match:
        return _money_token_to_number(digio_ticket_match.group(1)), default[1]

    single_ticket_k_match = re.search(r"\bTKT\s+(\d+(?:[,.]\d+)?)\s*K\b", normalized)
    if single_ticket_k_match:
        return _money_token_to_number(single_ticket_k_match.group(1), has_k=True), default[1]

    return default


def _money_token_to_number(value: str, has_k: bool = False) -> int:
    number = float(value.replace(",", "."))
    if has_k:
        number *= 1000
    return int(round(number))


def _normalize_text(value: object) -> str:
    normalized = _clean_text(value).upper()
    replacements = {
        "Á": "A",
        "À": "A",
        "Â": "A",
        "Ã": "A",
        "É": "E",
        "Ê": "E",
        "Í": "I",
        "Ó": "O",
        "Ô": "O",
        "Õ": "O",
        "Ú": "U",
        "Ç": "C",
    }
    for old, new in replacements.items():
        normalized = normalized.replace(old, new)
    return normalized


def _as_float(value: object) -> float:
    return float(str(value).replace(",", "."))


def _format_rate(value: float) -> str:
    return _format_percentage(value, places=2)


def _format_percentage(value: float, places: int) -> str:
    return f"{value:.{places}f}".replace(".", ",")


def _parse_date(value: object) -> datetime:
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass
    parsed = pd.to_datetime(text, dayfirst=True, errors="raise")
    return parsed.to_pydatetime()


def _write_workbook(rows: list[list[object | None]], output_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "RelatorioProdutos"
    sheet.sheet_view.showGridLines = False

    sheet.append(OUTPUT_HEADERS)
    for row in rows:
        sheet.append(row)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="Arial", size=10, color="000000")
    border = Border(
        left=Side(style="thin", color="D9E2F3"),
        right=Side(style="thin", color="D9E2F3"),
        top=Side(style="thin", color="D9E2F3"),
        bottom=Side(style="thin", color="D9E2F3"),
    )

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical="center")
            cell.border = border

    for row_idx in range(2, sheet.max_row + 1):
        sheet.cell(row_idx, 8).number_format = "yyyy-mm-dd"
        for col_idx in (10, 11, 17, 18, 19, 20):
            sheet.cell(row_idx, col_idx).number_format = "0"
        for col_idx in (22, 23, 24):
            sheet.cell(row_idx, col_idx).number_format = "0.00"
        sheet.cell(row_idx, 27).number_format = "0.0000"

    widths = {
        "A": 18,
        "B": 20,
        "C": 12,
        "D": 10,
        "E": 72,
        "H": 12,
        "J": 12,
        "K": 12,
        "L": 16,
        "M": 18,
        "U": 22,
        "X": 16,
        "AA": 18,
    }
    for col in range(1, sheet.max_column + 1):
        letter = sheet.cell(1, col).column_letter
        sheet.column_dimensions[letter].width = widths.get(letter, 14)

    sheet.freeze_panes = "A2"
    table_ref = f"A1:AO{sheet.max_row}"
    table = Table(displayName="Tabela2TechC6", ref=table_ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)

    workbook.save(output_path)
